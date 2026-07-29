from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from ccs_core import db_connection, record_audit

PRIVACY_LEVELS = ("public", "internal", "confidential")
APPROVAL_STATUSES = ("draft", "approved", "rejected")
PROVIDER_NAME = "local-evidence"
MAX_CHUNK_CHARS = 1200
CHUNK_OVERLAP = 180


@dataclass(frozen=True)
class Evidence:
    source_type: str
    source_id: int
    title: str
    source: str
    content: str
    score: int
    privacy_level: str


@dataclass(frozen=True)
class AssistantResponse:
    provider: str
    answer: str
    evidence: list[Evidence]
    privacy_level: str
    run_id: int


class AssistantProvider(Protocol):
    name: str
    allowed_privacy_levels: tuple[str, ...]

    def generate(self, question: str, evidence: list[Evidence]) -> str:
        ...


class LocalEvidenceProvider:
    name = PROVIDER_NAME
    allowed_privacy_levels = PRIVACY_LEVELS

    def generate(self, question: str, evidence: list[Evidence]) -> str:
        if not evidence:
            return (
                "Für diese Anfrage wurde keine ausreichend belastbare, freigegebene Quelle "
                "gefunden. Bitte den Fall fachlich prüfen, als Ticket dokumentieren und die "
                "Wissensbasis gezielt ergänzen."
            )

        steps: list[str] = []
        for item in evidence[:4]:
            excerpt = _compact_excerpt(item.content, 430)
            steps.append(f"- **{item.title}**: {excerpt}")

        return (
            "Auf Grundlage der freigegebenen internen Quellen ergibt sich folgender "
            "prüfpflichtiger Antwortentwurf:\n\n"
            + "\n".join(steps)
            + "\n\n**Kontrollhinweis:** Vor Versand fachlich prüfen. Der Entwurf führt "
              "keine Aktionen aus und ersetzt keine technische oder organisatorische Freigabe."
        )


def initialize_knowledge_ai() -> None:
    """Create additive schema extensions without invalidating an existing MVP database."""
    with db_connection() as db:
        article_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(knowledge_articles)").fetchall()
        }
        if "approval_status" not in article_columns:
            db.execute(
                "ALTER TABLE knowledge_articles "
                "ADD COLUMN approval_status TEXT NOT NULL DEFAULT 'approved'"
            )
        if "privacy_level" not in article_columns:
            db.execute(
                "ALTER TABLE knowledge_articles "
                "ADD COLUMN privacy_level TEXT NOT NULL DEFAULT 'internal'"
            )
        if "updated_at" not in article_columns:
            db.execute("ALTER TABLE knowledge_articles ADD COLUMN updated_at TEXT")

        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                filename TEXT NOT NULL,
                category TEXT NOT NULL,
                source TEXT,
                file_type TEXT NOT NULL,
                sha256 TEXT NOT NULL UNIQUE,
                approval_status TEXT NOT NULL
                    CHECK(approval_status IN ('draft', 'approved', 'rejected')),
                privacy_level TEXT NOT NULL
                    CHECK(privacy_level IN ('public', 'internal', 'confidential')),
                imported_by TEXT NOT NULL,
                imported_at TEXT NOT NULL,
                reviewed_by TEXT,
                reviewed_at TEXT
            );

            CREATE TABLE IF NOT EXISTS document_chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                chunk_no INTEGER NOT NULL,
                content TEXT NOT NULL,
                FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE,
                UNIQUE(document_id, chunk_no)
            );

            CREATE TABLE IF NOT EXISTS assistant_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                actor TEXT NOT NULL,
                provider TEXT NOT NULL,
                privacy_level TEXT NOT NULL,
                question TEXT NOT NULL,
                answer TEXT NOT NULL,
                evidence_refs TEXT,
                created_at TEXT NOT NULL
            );
            """
        )


def add_governed_article(
    *,
    title: str,
    category: str,
    content: str,
    source: str,
    approval_status: str,
    privacy_level: str,
    actor: str,
) -> int:
    _validate_status(approval_status)
    _validate_privacy(privacy_level)
    if not title.strip() or not content.strip():
        raise ValueError("Titel und Inhalt sind Pflichtfelder.")

    with db_connection() as db:
        cursor = db.execute(
            """
            INSERT INTO knowledge_articles (
                title, category, content, source, created_by, created_at,
                approval_status, privacy_level, updated_at
            ) VALUES (?, ?, ?, ?, ?, datetime('now'), ?, ?, datetime('now'))
            """,
            (
                title.strip(),
                category.strip() or "Allgemein",
                content.strip(),
                source.strip() or None,
                actor,
                approval_status,
                privacy_level,
            ),
        )
        article_id = int(cursor.lastrowid)

    record_audit(
        actor,
        "CREATE",
        "knowledge_article",
        str(article_id),
        f"Status: {approval_status}; Datenschutz: {privacy_level}; Titel: {title.strip()}",
    )
    return article_id


def set_article_status(article_id: int, status: str, actor: str) -> None:
    _validate_status(status)
    with db_connection() as db:
        updated = db.execute(
            """
            UPDATE knowledge_articles
            SET approval_status = ?, updated_at = datetime('now')
            WHERE id = ?
            """,
            (status, article_id),
        ).rowcount
    if updated == 0:
        raise ValueError("Wissensartikel wurde nicht gefunden.")
    record_audit(actor, "REVIEW", "knowledge_article", str(article_id), f"Status: {status}")


def list_governed_articles(include_unapproved: bool = True) -> list[dict]:
    initialize_knowledge_ai()
    where = "" if include_unapproved else "WHERE approval_status = 'approved'"
    with db_connection() as db:
        rows = db.execute(
            f"""
            SELECT id, title, category, content, source, created_by, created_at,
                   approval_status, privacy_level, updated_at
            FROM knowledge_articles
            {where}
            ORDER BY COALESCE(updated_at, created_at) DESC, id DESC
            """
        ).fetchall()
    return [dict(row) for row in rows]


def search_governed_knowledge(
    query: str,
    privacy_level: str = "internal",
    include_unapproved: bool = False,
) -> list[dict]:
    _validate_privacy(privacy_level)
    allowed = _allowed_privacy_levels(privacy_level)
    articles = list_governed_articles(include_unapproved=include_unapproved)
    filtered = [
        article
        for article in articles
        if article["privacy_level"] in allowed
        and (include_unapproved or article["approval_status"] == "approved")
    ]
    terms = _terms(query)
    if not terms:
        return filtered

    ranked = [(_score_text(article["title"], article["content"], terms), article) for article in filtered]
    return [
        article
        for score, article in sorted(ranked, key=lambda item: item[0], reverse=True)
        if score > 0
    ]


def import_document(
    *,
    filename: str,
    data: bytes,
    category: str,
    source: str,
    privacy_level: str,
    actor: str,
) -> int:
    initialize_knowledge_ai()
    _validate_privacy(privacy_level)
    if not filename.strip():
        raise ValueError("Dateiname fehlt.")
    if not data:
        raise ValueError("Die Datei ist leer.")

    file_type = Path(filename).suffix.lower().lstrip(".")
    if file_type not in {"txt", "pdf", "docx"}:
        raise ValueError("Unterstützt werden TXT, PDF und DOCX.")

    digest = hashlib.sha256(data).hexdigest()
    text = extract_document_text(filename, data)
    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("Aus dem Dokument konnte kein verwertbarer Text extrahiert werden.")

    with db_connection() as db:
        duplicate = db.execute(
            "SELECT id FROM documents WHERE sha256 = ?",
            (digest,),
        ).fetchone()
        if duplicate:
            raise ValueError(f"Dieses Dokument ist bereits als ID {duplicate['id']} vorhanden.")

        cursor = db.execute(
            """
            INSERT INTO documents (
                filename, category, source, file_type, sha256, approval_status,
                privacy_level, imported_by, imported_at
            ) VALUES (?, ?, ?, ?, ?, 'draft', ?, ?, datetime('now'))
            """,
            (
                filename.strip(),
                category.strip() or "Support",
                source.strip() or None,
                file_type,
                digest,
                privacy_level,
                actor,
            ),
        )
        document_id = int(cursor.lastrowid)
        db.executemany(
            """
            INSERT INTO document_chunks (document_id, chunk_no, content)
            VALUES (?, ?, ?)
            """,
            [(document_id, index, chunk) for index, chunk in enumerate(chunks, start=1)],
        )

    record_audit(
        actor,
        "IMPORT",
        "document",
        str(document_id),
        f"{filename}; Chunks: {len(chunks)}; Status: draft; Datenschutz: {privacy_level}",
    )
    return document_id


def extract_document_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
            try:
                return _normalize_text(data.decode(encoding))
            except UnicodeDecodeError:
                continue
        raise ValueError("TXT-Datei konnte nicht dekodiert werden.")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF-Import benötigt das Paket 'pypdf'.") from exc
        reader = PdfReader(io.BytesIO(data))
        return _normalize_text("\n".join((page.extract_text() or "") for page in reader.pages))

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX-Import benötigt das Paket 'python-docx'.") from exc
        document = Document(io.BytesIO(data))
        return _normalize_text("\n".join(paragraph.text for paragraph in document.paragraphs))

    raise ValueError("Nicht unterstütztes Dateiformat.")


def chunk_text(text: str, max_chars: int = MAX_CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    normalized = _normalize_text(text)
    if not normalized:
        return []
    if max_chars < 200:
        raise ValueError("Chunk-Größe muss mindestens 200 Zeichen betragen.")
    if overlap < 0 or overlap >= max_chars:
        raise ValueError("Überlappung muss zwischen 0 und Chunk-Größe liegen.")

    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + max_chars, len(normalized))
        if end < len(normalized):
            break_at = normalized.rfind(" ", start + max_chars // 2, end)
            if break_at > start:
                end = break_at
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(end - overlap, start + 1)
    return chunks


def list_documents() -> list[dict]:
    initialize_knowledge_ai()
    with db_connection() as db:
        rows = db.execute(
            """
            SELECT d.id, d.filename, d.category, d.source, d.file_type,
                   d.approval_status, d.privacy_level, d.imported_by,
                   d.imported_at, d.reviewed_by, d.reviewed_at,
                   COUNT(c.id) AS chunk_count
            FROM documents d
            LEFT JOIN document_chunks c ON c.document_id = d.id
            GROUP BY d.id
            ORDER BY d.id DESC
            """
        ).fetchall()
    return [dict(row) for row in rows]


def set_document_status(document_id: int, status: str, actor: str) -> None:
    _validate_status(status)
    with db_connection() as db:
        updated = db.execute(
            """
            UPDATE documents
            SET approval_status = ?, reviewed_by = ?, reviewed_at = datetime('now')
            WHERE id = ?
            """,
            (status, actor, document_id),
        ).rowcount
    if updated == 0:
        raise ValueError("Dokument wurde nicht gefunden.")
    record_audit(actor, "REVIEW", "document", str(document_id), f"Status: {status}")


def retrieve_evidence(
    question: str,
    privacy_level: str = "internal",
    limit: int = 5,
) -> list[Evidence]:
    initialize_knowledge_ai()
    _validate_privacy(privacy_level)
    terms = _terms(question)
    if not terms:
        return []
    allowed = _allowed_privacy_levels(privacy_level)
    placeholders = ",".join("?" for _ in allowed)

    evidence: list[Evidence] = []
    with db_connection() as db:
        articles = db.execute(
            f"""
            SELECT id, title, content, COALESCE(source, 'Wissensartikel') AS source,
                   privacy_level
            FROM knowledge_articles
            WHERE approval_status = 'approved'
              AND privacy_level IN ({placeholders})
            """,
            allowed,
        ).fetchall()
        chunks = db.execute(
            f"""
            SELECT c.id, c.content, d.filename AS title,
                   COALESCE(d.source, d.filename) AS source, d.privacy_level
            FROM document_chunks c
            JOIN documents d ON d.id = c.document_id
            WHERE d.approval_status = 'approved'
              AND d.privacy_level IN ({placeholders})
            """,
            allowed,
        ).fetchall()

    for row in articles:
        score = _score_text(row["title"], row["content"], terms)
        if score:
            evidence.append(
                Evidence(
                    source_type="article",
                    source_id=int(row["id"]),
                    title=row["title"],
                    source=row["source"],
                    content=row["content"],
                    score=score,
                    privacy_level=row["privacy_level"],
                )
            )

    for row in chunks:
        score = _score_text(row["title"], row["content"], terms)
        if score:
            evidence.append(
                Evidence(
                    source_type="document_chunk",
                    source_id=int(row["id"]),
                    title=row["title"],
                    source=row["source"],
                    content=row["content"],
                    score=score,
                    privacy_level=row["privacy_level"],
                )
            )

    return sorted(evidence, key=lambda item: item.score, reverse=True)[:limit]


def generate_assistant_answer(
    *,
    question: str,
    privacy_level: str,
    actor: str,
    provider_name: str = PROVIDER_NAME,
) -> AssistantResponse:
    initialize_knowledge_ai()
    _validate_privacy(privacy_level)
    if not question.strip():
        raise ValueError("Bitte eine Supportfrage eingeben.")

    provider = get_provider(provider_name)
    if privacy_level not in provider.allowed_privacy_levels:
        raise ValueError(
            f"Provider '{provider.name}' ist für Datenschutzstufe '{privacy_level}' nicht freigegeben."
        )

    evidence = retrieve_evidence(question, privacy_level=privacy_level)
    answer = provider.generate(question.strip(), evidence)
    references = ";".join(f"{item.source_type}:{item.source_id}" for item in evidence)

    with db_connection() as db:
        cursor = db.execute(
            """
            INSERT INTO assistant_runs (
                actor, provider, privacy_level, question, answer,
                evidence_refs, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, datetime('now'))
            """,
            (
                actor,
                provider.name,
                privacy_level,
                question.strip(),
                answer,
                references or None,
            ),
        )
        run_id = int(cursor.lastrowid)

    record_audit(
        actor,
        "GENERATE",
        "assistant_run",
        str(run_id),
        f"Provider: {provider.name}; Datenschutz: {privacy_level}; Quellen: {len(evidence)}",
    )
    return AssistantResponse(
        provider=provider.name,
        answer=answer,
        evidence=evidence,
        privacy_level=privacy_level,
        run_id=run_id,
    )


def list_assistant_runs(limit: int = 100) -> list[dict]:
    initialize_knowledge_ai()
    with db_connection() as db:
        rows = db.execute(
            """
            SELECT id, actor, provider, privacy_level, question, answer,
                   evidence_refs, created_at
            FROM assistant_runs
            ORDER BY id DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()
    return [dict(row) for row in rows]


def get_provider(name: str) -> AssistantProvider:
    normalized = name.strip().lower()
    if normalized == PROVIDER_NAME:
        return LocalEvidenceProvider()
    raise ValueError(
        "Unbekannter oder nicht freigegebener Provider. "
        "In Version 0.2.0 ist ausschließlich 'local-evidence' aktiv."
    )


def _allowed_privacy_levels(requested: str) -> tuple[str, ...]:
    rank = PRIVACY_LEVELS.index(requested)
    return PRIVACY_LEVELS[: rank + 1]


def _validate_privacy(value: str) -> None:
    if value not in PRIVACY_LEVELS:
        raise ValueError(f"Ungültige Datenschutzstufe: {value}")


def _validate_status(value: str) -> None:
    if value not in APPROVAL_STATUSES:
        raise ValueError(f"Ungültiger Freigabestatus: {value}")


def _terms(text: str) -> list[str]:
    return [
        token.lower()
        for token in re.findall(r"[\wäöüÄÖÜß-]+", text, flags=re.UNICODE)
        if len(token) >= 3
    ]


def _score_text(title: str, content: str, terms: list[str]) -> int:
    title_lower = title.lower()
    content_lower = content.lower()
    score = 0
    for term in terms:
        score += title_lower.count(term) * 5
        score += content_lower.count(term)
    return score


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    return re.sub(r"\s+", " ", text).strip()


def _compact_excerpt(text: str, limit: int) -> str:
    normalized = _normalize_text(text)
    if len(normalized) <= limit:
        return normalized
    return normalized[: limit - 1].rstrip() + "…"
